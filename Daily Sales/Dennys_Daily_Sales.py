"""PrimeLedgerAI.com - Denny's Daily Sales to QuickBooks CSV converter."""

from __future__ import annotations

import csv
import json
import os
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("Missing openpyxl. Run: py -m pip install openpyxl") from exc


CENT = Decimal("0.01")

# Exact QuickBooks locations and cash accounts supplied by Marietta Hotel.
# The store number is used first so variations such as "5164: Berkshire" still map correctly.
LOCATION_AND_CASH = {
    "5164": ("5164-Berkshire OH", "Cash on Hand - Berkshire OH"),
    "5165": ("5165-Perrysburg OH", "Cash on Hand - Perrysburg,OH"),
    "5166": ("5166-Jeffersonville", "Cash on Hand - Jeffersonville"),
    "5167": ("5167-Elizabethtown KY", "Cash on Hand - Elizabethtown-KY"),
    "5168": ("5168-Catlettsburg KY", "Cash on hand - Catlettsburg KY"),
    "5169": ("5169-Columbus OH", "Cash on Hand - Columbus OH"),
    "7401": ("7401-Youngstown", "Cash on Hand - Youngstown"),
    "9690": ("9690 - Marietta Rest", "Cash on Hand - Marietta"),
    "9697": ("9697-Findley", "Cash on Hand - Findlay"),
}

NAMED_LOCATION_AND_CASH = {
    "marietta hotel": ("Marietta Hotel", "Cash on Hand - Marietta Hotel"),
}


SETTINGS_PATH = Path(os.environ.get("APPDATA") or str(Path.home() / ".config")) / "PrimeLedgerAI" / "daily_sales_mapping.json"

BASE_MAPPING = {
    "Net Sales": "Sales A/C",
    "Total Sales Tax": "Tax Payable:Sales Tax Payable",
    "Total Discounts": "Food Discount A/c",
    "Credit Tips Received": "Tips Payable",
    "Credit Tips Paid": "Tips Payable",
    "Cash Paid Outs - Non Adjusting": "Cash Paid Outs",
    "Cash Paid Outs - Sales Adjusting": "Cash Paid Outs",
    "Auto-balance: source tender variance": "Daily Sales Variance",
}


def default_mapping():
    result = dict(BASE_MAPPING)
    for location, cash in list(LOCATION_AND_CASH.values()) + list(NAMED_LOCATION_AND_CASH.values()):
        result["Cash | " + location] = cash
    for name in ("Door Dash", "Doordash Banda Burrito", "Doordash VRM", "Doordash Nathans",
                 "Uber Eats", "Uber Eats Banda Burrito", "Uber Eats VR", "Uber Eats VRM", "Uber Eats Nathans",
                 "Online", "Grubhub", "Grub Hub", "Grubhub Banda Burrito", "Grubhub VRB", "Grubhub VRM", "QR Pay", "Master Card", "Visa", "Discover", "Debit",
                 "American Express", "Gift Card"):
        result["Tender: " + name] = tender_account("Tender: " + name)
    return result


def load_settings(path=SETTINGS_PATH):
    if not path.exists():
        return {"accounts": [], "mapping": default_mapping()}
    data = json.loads(path.read_text(encoding="utf-8"))
    if (not isinstance(data, dict) or not isinstance(data.get("accounts"), list)
            or not all(isinstance(x, str) for x in data["accounts"])
            or not isinstance(data.get("mapping"), dict)
            or not all(isinstance(k, str) and isinstance(v, str) for k, v in data["mapping"].items())):
        raise ValueError("The saved account mapping is invalid. Restore your mapping backup.")
    return data


def save_settings(data, path=SETTINGS_PATH):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    temporary.replace(path)


def read_accounts(path):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        try:
            content = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            content = path.read_text(encoding="cp1252")
        rows = list(csv.reader(content.splitlines()))
    else:
        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            rows = list(wb.active.iter_rows(values_only=True))
        finally:
            wb.close()
    for index, row in enumerate(rows[:30]):
        headers = [re.sub(r"[^a-z]", "", str(x or "").lower()) for x in row]
        for preferred in ("fullaccountname", "accountfullname", "accountname", "account"):
            if preferred in headers:
                column = headers.index(preferred)
                names = sorted({str(r[column]).strip() for r in rows[index+1:]
                                if len(r) > column and r[column] is not None and str(r[column]).strip()}, key=str.casefold)
                if not names:
                    raise ValueError("No accounts found in this chart of accounts.")
                return names
    raise ValueError("Chart of accounts needs an Account name or Full account name column. Include Parent:Subaccount names exactly as required by your importer.")


def mapping_key(description, location):
    if description == "Tender: Cash (net)":
        return "Cash | " + location
    if description == "Total Discounts - Offset":
        return "Total Discounts"
    return description


def money(value) -> Decimal:
    if value is None or str(value).strip() == "":
        return Decimal("0.00")
    try:
        return Decimal(str(value).replace("$", "").replace(",", "").strip()).quantize(CENT)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid dollar amount: {value!r}") from exc


def read_source(path: Path):
    rows = []
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))
    else:
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(x or "").strip() for x in values[0]]
        rows = [dict(zip(headers, row)) for row in values[1:]]

    required = {"Date", "Store", "Field", "Value"}
    if not rows and path.stat().st_size:
        raise ValueError("The input file has no data rows.")
    if rows and not required.issubset(rows[0]):
        raise ValueError("Required columns are: Date, Store, Field, Value")
    return rows


def parse_date(raw):
    text = str(raw).strip()
    if text.endswith(".0"):
        text = text[:-2]
    for fmt in ("%Y%m%d", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"Unsupported date: {raw!r}")


def location_details(store: str):
    raw = str(store).strip()
    match = re.match(r"\s*(\d+)", raw)
    if match and match.group(1) in LOCATION_AND_CASH:
        return LOCATION_AND_CASH[match.group(1)]
    named = NAMED_LOCATION_AND_CASH.get(raw.casefold())
    if named:
        return named
    valid = ", ".join(location for location, _ in LOCATION_AND_CASH.values())
    raise ValueError(f"Store {raw!r} is not in the location mapping. Valid sales locations: {valid}, Marietta Hotel")


def tender_account(field: str):
    name = field.removeprefix("Tender:").strip()
    low = name.lower()
    # Match the specific DoorDash brand before the regular DoorDash tender.
    # Account spelling is copied from the supplied QuickBooks chart of accounts.
    compact = re.sub(r"[^a-z0-9]", "", low)
    if "doordash" in compact:
        if "amendment" in compact:
            return "Credit Sales:Amendments - Doordash"
        if "vrbur" in compact or "bandaburrito" in compact:
            return "Credit Sales:DoorDash - VRBUR"
        if "vrnat" in compact or "nathan" in compact:
            return "Credit Sales:Doordash - VRNAT"
        if "vrm" in compact:
            return "Credit Sales:Doordash - VRM"
        if "vrb" in compact:
            return "Credit Sales:DoorDash - VRB"
        if compact not in {"doordash", "doordashsales"}:
            raise ValueError(f"DoorDash tender needs an account mapping: {name!r}")
        return "Credit Sales:Doordash - Sales"
    # User confirmed that Online payments belong to DoorDash.
    if low == "online":
        return "Credit Sales:Doordash - Sales"
    if "uber" in low:
        if "vrbur" in compact or "bandaburrito" in compact:
            return "Credit Sales:Uber Eats - VRBUR"
        if "vrnat" in compact or "nathan" in compact:
            return "Credit Sales:UberEat- VRNAT"
        if "vrm" in compact:
            return "Credit Sales:Uber Eats - VRM"
        if compact.endswith("vr"):
            return "Credit Sales:UberEat - VR"
        if compact not in {"uber", "ubereat", "ubereats", "ubereatsales", "ubereatssales"}:
            raise ValueError(f"Uber tender needs an account mapping: {name!r}")
        return "Credit Sales:Uber Eats - Sales"
    if "grub" in low:
        # Brand accounts first; "vrb" is a prefix of "vrbur", so test VRBUR before VRB.
        if "vrbur" in compact or "bandaburrito" in compact:
            return "Credit Sales:Grubhub - VRBUR"
        if "vrnat" in compact or "nathan" in compact:
            return "Credit Sales:Grubhub - VRNAT"
        if "vrm" in compact:
            return "Credit Sales:Grubhub - VRM"
        if "vrb" in compact:
            return "Credit Sales:Grubhub- VRB"
        return "Credit Sales:Grubhub - Sales"
    if low == "qr pay":
        return "Credit Sales:QR Pay - Sales"
    if low in {"master card", "visa", "discover", "debit", "american express"}:
        return "Credit Sales:Credit card - Sales to Dennys"
    if low == "gift card":
        return "Gift Card Payable"
    if low == "cash":
        return "Cash on Hand"
    return f"Credit Sales:{name}"


def add_line(lines, account, amount, side, description, location):
    amount = money(amount)
    if amount == 0:
        return
    # Negative source values automatically reverse debit/credit.
    if amount < 0:
        side = "credit" if side == "debit" else "debit"
        amount = -amount
    lines.append({
        "Account": account,
        "Debits": amount if side == "debit" else None,
        "Credits": amount if side == "credit" else None,
        "Description": description,
        "Name": "",
        "Location": location,
    })


def build_journal(field_values, location, cash_account, settings=None):
    lines = []

    def add_line(lines, account, amount, side, description, location):
        if money(amount) == 0:
            return
        if settings is not None:
            key = mapping_key(description, location)
            account = settings["mapping"].get(key, "")
            if not account or account not in settings["accounts"]:
                raise ValueError(f"Open Account Mapping and select an account from your chart for: {key}")
        globals()["add_line"](lines, account, amount, side, description, location)

    get = lambda key: field_values.get(key, Decimal("0.00"))

    add_line(lines, "Sales A/C", get("Net Sales"), "credit", "Net Sales", location)
    add_line(lines, "Tax Payable:Sales Tax Payable", get("Total Sales Tax"), "credit", "Total Sales Tax", location)

    discount = get("Total Discounts")
    add_line(lines, "Food Discount A/c", discount, "debit", "Total Discounts", location)
    add_line(lines, "Food Discount A/c", discount, "credit", "Total Discounts - Offset", location)

    tips_received = get("Credit Tips Received")
    tips_paid = get("Credit Tips Paid")
    add_line(lines, "Tips Payable", tips_received, "credit", "Credit Tips Received", location)
    add_line(lines, "Tips Payable", tips_paid, "debit", "Credit Tips Paid", location)

    paid_out_non = get("Cash Paid Outs - Non Adjusting")
    paid_out_sales = get("Cash Paid Outs - Sales Adjusting")
    add_line(lines, "Cash Paid Outs", paid_out_non, "debit", "Cash Paid Outs - Non Adjusting", location)
    add_line(lines, "Cash Paid Outs", paid_out_sales, "debit", "Cash Paid Outs - Sales Adjusting", location)

    for field, amount in field_values.items():
        if not field.startswith("Tender:") or amount == 0:
            continue
        tender = field.removeprefix("Tender:").strip()
        if tender.lower() == "cash":
            # Cash tender includes tips. Tips paid and paid-outs reduce actual cash retained.
            cash_net = amount - tips_paid - paid_out_non - paid_out_sales
            add_line(lines, cash_account, cash_net, "debit", "Tender: Cash (net)", location)
        else:
            add_line(lines, tender_account(field) if settings is None else "", amount, "debit", field, location)

    # Keep empty/all-zero journals in Decimal arithmetic: sum([]) is an int.
    debits = sum((x["Debits"] or Decimal("0") for x in lines), Decimal("0.00"))
    credits = sum((x["Credits"] or Decimal("0") for x in lines), Decimal("0.00"))
    variance = (debits - credits).quantize(CENT, rounding=ROUND_HALF_UP)
    if variance > 0:
        add_line(lines, "Daily Sales Variance", variance, "credit", "Auto-balance: source tender variance", location)
    elif variance < 0:
        add_line(lines, "Daily Sales Variance", -variance, "debit", "Auto-balance: source tender variance", location)
    return lines, variance


def convert(input_path: str, output_path: str, settings=None):
    if Path(input_path).resolve() == Path(output_path).resolve():
        raise ValueError("Choose an output file different from the input file.")
    source = read_source(Path(input_path))
    grouped = defaultdict(dict)
    for row in source:
        if not str(row.get("Store") or "").strip() or not str(row.get("Field") or "").strip():
            continue
        key = (parse_date(row["Date"]), str(row["Store"]).strip())
        grouped[key][str(row["Field"]).strip()] = money(row.get("Value"))
    if not grouped:
        raise ValueError("No usable daily-sales records were found.")

    headers = ["Journal No.", "Journal Date", "Account", "Debits", "Credits", "Description", "Name", "Location"]
    output_rows = []

    for (date, store), fields in sorted(grouped.items()):
        location, cash_account = location_details(store)
        store_match = re.match(r"\d+", location)
        store_no = store_match.group(0) if store_match else re.sub(r"\W+", "", location)[:8]
        journal_no = f"DS{date:%m%d}-{store_no}"
        lines, variance = build_journal(fields, location, cash_account, settings)
        for line in lines:
            output_rows.append({
                "Journal No.": journal_no,
                "Journal Date": date.strftime("%m/%d/%Y"),
                "Account": line["Account"],
                "Debits": f'{line["Debits"]:.2f}' if line["Debits"] is not None else "",
                "Credits": f'{line["Credits"]:.2f}' if line["Credits"] is not None else "",
                "Description": line["Description"],
                "Name": line["Name"],
                "Location": line["Location"],
            })

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    with Path(output_path).open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(output_rows)
    return len(grouped), len(output_rows)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("PrimeLedgerAI - Denny's Daily Sales Converter")
        self.geometry("800x390")
        self.resizable(False, False)
        self.input_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.settings = load_settings()
        self.mapping_status = tk.StringVar(value=f"{len(self.settings['accounts'])} saved accounts — mapping reused automatically")

        ttk.Label(self, text="Denny's Daily Sales → CSV Journal", font=("Segoe UI", 18, "bold")).pack(pady=(20, 18))
        form = ttk.Frame(self, padding=12)
        form.pack(fill="x")
        ttk.Label(form, text="Input CSV/Excel:").grid(row=0, column=0, sticky="w", pady=8)
        ttk.Entry(form, textvariable=self.input_var, width=68).grid(row=0, column=1, padx=8)
        ttk.Button(form, text="Browse", command=self.choose_input).grid(row=0, column=2)
        ttk.Label(form, text="Output CSV:").grid(row=1, column=0, sticky="w", pady=8)
        ttk.Entry(form, textvariable=self.output_var, width=68).grid(row=1, column=1, padx=8)
        ttk.Button(form, text="Browse", command=self.choose_output).grid(row=1, column=2)
        ttk.Button(self, text="ACCOUNT MAPPING / LOAD CHART OF ACCOUNTS", command=self.edit_mapping).pack(pady=5)
        ttk.Label(self, textvariable=self.mapping_status).pack()
        ttk.Label(self, text="First time: select your sales input, load your chart, review mappings and save.").pack(pady=5)
        ttk.Button(self, text="CREATE CSV JOURNAL", command=self.run, width=30).pack(pady=20)

    def edit_mapping(self):
        mapping = dict(self.settings["mapping"])
        accounts = list(self.settings["accounts"])
        if self.input_var.get():
            try:
                for row in read_source(Path(self.input_var.get())):
                    field = str(row.get("Field") or "").strip()
                    if field.startswith("Tender:") and field.removeprefix("Tender:").strip().lower() != "cash":
                        if field not in mapping:
                            try:
                                mapping[field] = tender_account(field)
                            except ValueError:
                                mapping[field] = ""
            except Exception as exc:
                messagebox.showerror("Could not read sales input", str(exc))
                return
        win = tk.Toplevel(self)
        win.title("Account Mapping — PrimeLedgerAI")
        win.geometry("1050x680")
        win.transient(self)
        win.grab_set()
        ttk.Label(win, text="1. Load your chart   2. Select a row and account   3. Apply   4. Save mappings",
                  font=("Segoe UI", 11, "bold")).pack(pady=12)
        ttk.Label(win, text="Load a sales input first to discover new tenders. Cash accounts are mapped separately by location.").pack()
        toolbar = ttk.Frame(win, padding=8)
        toolbar.pack(fill="x")
        status = tk.StringVar()
        frame = ttk.Frame(win, padding=8)
        frame.pack(fill="both", expand=True)
        tree = ttk.Treeview(frame, columns=("field", "account", "status"), show="headings", selectmode="browse")
        for col, title, width in (("field", "Sales field / location", 300), ("account", "QuickBooks account", 480), ("status", "Status", 130)):
            tree.heading(col, text=title)
            tree.column(col, width=width)
        scrollbar = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side="right", fill="y")
        tree.pack(fill="both", expand=True)
        editor = ttk.Frame(win, padding=10)
        editor.pack(fill="x")
        selected = tk.StringVar(value="Select a mapping row above")
        ttk.Label(editor, textvariable=selected).pack(anchor="w")
        account_var = tk.StringVar()
        combo = ttk.Combobox(editor, textvariable=account_var, values=accounts, width=105)
        combo.pack(fill="x", pady=6)
        ttk.Label(editor, text="Type to filter account names, then choose the exact account from the list.").pack(anchor="w")

        def refresh():
            current = tree.selection()
            tree.delete(*tree.get_children())
            for key, account in sorted(mapping.items()):
                tree.insert("", "end", iid=key, values=(key, account, "Ready" if account in accounts else "Select account"))
            if current and current[0] in mapping:
                tree.selection_set(current[0])
            status.set(f"{len(accounts)} accounts loaded. Only accounts in your chart can be used.")
            combo.configure(values=accounts)

        def load_chart():
            path = filedialog.askopenfilename(parent=win, filetypes=[("Chart of accounts", "*.csv *.xlsx *.xlsm")])
            if path:
                try:
                    loaded = read_accounts(path)
                    accounts[:] = loaded
                    refresh()
                except Exception as exc:
                    messagebox.showerror("Could not load chart", str(exc), parent=win)

        def select_row(event=None):
            if tree.selection():
                key = tree.selection()[0]
                selected.set(key)
                account_var.set(mapping[key])
                combo.configure(values=accounts)

        def apply():
            if not tree.selection():
                messagebox.showerror("Select row", "Select the sales field you want to map.", parent=win)
                return False
            account = account_var.get().strip()
            if account not in accounts:
                messagebox.showerror("Select account", "Choose an exact account from your loaded chart.", parent=win)
                return False
            mapping[tree.selection()[0]] = account
            refresh()
            return True

        def save():
            if not accounts:
                messagebox.showerror("Load chart", "Load your chart of accounts first.", parent=win)
                return
            if tree.selection() and account_var.get().strip() != mapping[tree.selection()[0]]:
                if not apply():
                    return
            data = {"accounts": accounts, "mapping": mapping}
            try:
                save_settings(data)
            except Exception as exc:
                messagebox.showerror("Could not save", str(exc), parent=win)
                return
            self.settings = data
            self.mapping_status.set(f"{len(accounts)} saved accounts — mapping reused automatically")
            messagebox.showinfo("Saved", "Your chart and mappings are saved for future use. You can change them here anytime.\n\nAny unmapped field used in a journal must be mapped before export.", parent=win)
            win.destroy()

        tree.bind("<<TreeviewSelect>>", select_row)
        combo.bind("<KeyRelease>", lambda event: combo.configure(values=[a for a in accounts if account_var.get().casefold() in a.casefold()]))
        ttk.Button(toolbar, text="LOAD / REPLACE CHART OF ACCOUNTS", command=load_chart).pack(side="left")
        ttk.Label(toolbar, textvariable=status).pack(side="left", padx=12)
        buttons = ttk.Frame(win, padding=10)
        buttons.pack(fill="x")
        ttk.Button(buttons, text="Apply selected account", command=apply).pack(side="left")
        ttk.Button(buttons, text="SAVE MAPPINGS", command=save).pack(side="right")
        ttk.Button(buttons, text="Cancel", command=win.destroy).pack(side="right", padx=10)
        refresh()

    def choose_input(self):
        path = filedialog.askopenfilename(filetypes=[("Daily sales files", "*.csv *.xlsx *.xlsm"), ("All files", "*.*")])
        if path:
            self.input_var.set(path)
            p = Path(path)
            self.output_var.set(str(p.with_name(p.stem + "_Journal_Output.csv")))

    def choose_output(self):
        path = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV file", "*.csv")])
        if path:
            self.output_var.set(path)

    def run(self):
        try:
            if not self.input_var.get() or not self.output_var.get():
                raise ValueError("Please select the input and output files.")
            if not self.settings["accounts"]:
                raise ValueError("Open Account Mapping and load your chart of accounts first.")
            journals, lines = convert(self.input_var.get(), self.output_var.get(), self.settings)
            messagebox.showinfo("Completed", f"Created {journals} journal entries and {lines} journal lines.\n\n{self.output_var.get()}")
        except Exception as exc:
            messagebox.showerror("Could not create output", str(exc))


if __name__ == "__main__":
    try:
        App().mainloop()
    except Exception as exc:
        messagebox.showerror("Could not start", str(exc))
