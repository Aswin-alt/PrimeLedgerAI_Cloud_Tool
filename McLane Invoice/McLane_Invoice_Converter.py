"""PrimeLedgerAI.com - McLane Invoice Detail to QuickBooks Online Bill CSV.

Simple desktop converter based on the supplied McLane invoice example.
- Select one McLane Invoice Detail XLSX or CSV.
- The program groups lines by invoice and CustGLDescription.
- It maps each group to the supplied QuickBooks Chart of Accounts.
- It derives the QuickBooks location from the McLane Unit/store number.
- It saves a QuickBooks-ready CSV using the exact 19-column layout from the sample.

For XLSX files this program requires openpyxl:
    py -m pip install openpyxl

Run:
    py McLane_Invoice_Converter.py
"""
from __future__ import annotations

import csv
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

CENT = Decimal("0.01")

OUTPUT_HEADERS = [
    "Bill Number", "Vendor", "Mailing Address", "Terms", "BillDate", "DueDate",
    "Location", "Memo", "Type", "Category/Account", "Product/Service", "Qty",
    "Rate", "Description", "Amount", "Billable", "Customer/Project", "TaxRate", "Class"
]

# Exact location names from the supplied QuickBooks Location List.
LOCATION_BY_STORE = {
    "5164": "5164-Berkshire OH",
    "5165": "5165-Perrysburg OH",
    "5166": "5166-Jeffersonville",
    "5167": "5167-Elizabethtown KY",
    "5168": "5168-Catlettsburg KY",
    "5169": "5169-Columbus OH",
    "7401": "7401-Youngstown",
    "9690": "9690 - Marietta Rest",
    "9697": "9697-Findley",
}

# Vendor names. Youngstown is confirmed by the supplied output example.
# Edit these values if your QuickBooks vendor names are different.
VENDOR_BY_STORE = {
    "5164": "McLane - Nights and Bites",
    "5165": "McLane - Nights and Bites",
    "5166": "McLane - Nights and Bites",
    "5167": "McLane - Nights and Bites",
    "5168": "McLane - Nights and Bites",
    "5169": "McLane - Nights and Bites",
    "7401": "McLane - Youngstown",
    "9690": "McLane - Marietta",
    "9697": "McLane - Nights and Bites",
}

# Exact GL Description -> coa mappings from MCLAN ACCOUNT MAPPING.xlsx.
# Keep QuickBooks account spelling exactly as supplied.
ACCOUNT_MAP = {
    'GLASSWARE': 'Non Ingredient - Goods',
    'BREAD & ROLLS': 'Ingredient cost:BREAD & ROLLS',
    'BUTTER': 'Ingredient cost:BUTTER',
    'CHEESE': 'Ingredient cost:CHEESE',
    'DAIRY': 'Ingredient cost:DAIRY',
    'DESSERT, OTHER INGRED': 'Ingredient cost:DESSERT OTHER INGRED',
    'EARLY PAY DISCOUNT': 'Ingredient cost:EARLY PAY DISCOUNT',
    'EGGS': 'Ingredient cost:EGGS',
    'FISH': 'Ingredient cost:FISH',
    'FRENCH FRIES': 'Ingredient cost:FRENCH FRIES',
    'GRAVY/SAUCE': 'Ingredient cost:GRAVY/SAUCE',
    'GROUND BEEF': 'Ingredient cost:GROUND BEEF',
    'HASH BROWNS': 'Ingredient cost:HASH BROWNS',
    'LIQUID SHORTENING': 'Non Ingredient - Goods:liquid',
    'ORANGE JUICE': 'Ingredient cost:ORANGE JUICE',
    'OTHER GROCERIES': 'Ingredient cost:OTHER GROCERIES',
    'OTHER MEATS': 'Ingredient cost:OTHER MEATS',
    'OTHER REPLACEMENTS': 'Non Ingredient - Goods:OTHER REPLACEMENTS',
    'PANCAKE MIX': 'Ingredient cost:PANCAKE MIX',
    'PAPER PRODUCTS': 'Non Ingredient - Goods:PAPER PRODUCTS',
    'PORK': 'Ingredient cost:PORK',
    'POULTRY': 'Ingredient cost:POULTRY',
    'PRODUCE': 'Ingredient cost:PRODUCE',
    'SOFT DRINKS': 'Ingredient cost:SOFT DRINKS',
    'STEAK': 'Ingredient cost:STEAK',
    'PIES': 'Ingredient cost:PIES',
    'OTHER': 'Ingredient cost:Other',
    'ICE CREAM': 'Ingredient cost:Ice Cream',
    'KIDS PROGRAM': 'Ingredient cost:KIDS PROGRAM',
    'MAPLE SYRUP': 'Ingredient cost:MAPLE SYRUP',
    'PRIME RIB': 'Ingredient cost:prime rib',
    'SLICED STRAWBERRIES': 'Ingredient cost:Sliced Strawberries',
    'SOUPS': 'Ingredient cost:SOUPS',
    'CHINA': 'Ingredient cost:china',
    'COFFEE': 'Ingredient cost:COFFEE',
    'CONDIMENTS': 'Ingredient cost:Condiments',
    'COST OF LABOR - COGS': 'Ingredient cost:Cost of labor - COGS',
    'EQUIPMENT RENTAL - COGS': 'Ingredient cost:Equipment rental - COGS',
    'CAKES & DESSERTS': 'Ingredient cost:CAKES & DESSERTS',
    'CATSUP': 'Ingredient cost:CATSUP',
    'FREIGHT IN - COGS': 'Ingredient cost:Freight in - COGS',
    'GOODS': 'Non Ingredient - Goods',
    'JANITORIAL GOODS': 'Non Ingredient - Goods:JANITORIAL GOODS',
    'CHEMICAL PURCHASES': 'Non Ingredient - Goods:Chemical Purchases',
    'DETERGENTS': 'Non Ingredient - Goods:DETERGENTS',
    'OIL': 'Non Ingredient - Goods:oil',
    'SILVERWARE': 'Non Ingredient - Goods:SILVERWARE',
    'SMALLWARES': 'Non Ingredient - Goods:SMALLWARES',
}


def normalize(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).upper()


def money(value) -> Decimal:
    text = str(value if value is not None else "").strip().replace("$", "").replace(",", "")
    if not text:
        return Decimal("0.00")
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid amount: {value}") from exc
    return amount.quantize(CENT, rounding=ROUND_HALF_UP)


def parse_date(value) -> datetime:
    if isinstance(value, datetime):
        return value
    text = str(value or "").strip()
    for fmt in (
        "%m/%d/%Y %I:%M:%S %p",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    raise ValueError(f"Invalid date: {value}")


# McLane old unit number -> current QuickBooks store/location number.
# McLane input may contain extra leading digits, so the converter first uses
# the last 4 digits as the old store number, then converts it to the new one.
OLD_TO_NEW_STORE = {
    "8590": "5164",  # Berkshire OH
    "8611": "5165",  # Perrysburg OH
    "8631": "5166",  # Jeffersonville
    "8768": "5167",  # Elizabethtown KY
    "8772": "5168",  # Catlettsburg KY
    "8829": "5169",  # Columbus OH
    "7401": "7401",  # Youngstown
    "9690": "9690",  # Marietta Rest
    "9697": "9697",  # Findlay
}


def store_number(unit) -> str:
    digits = re.sub(r"\D", "", str(unit or ""))
    if len(digits) < 4:
        raise ValueError(f"Invalid McLane Unit: {unit}")

    old_store = digits[-4:]
    new_store = OLD_TO_NEW_STORE.get(old_store)
    if not new_store:
        raise ValueError(
            f"McLane Unit {unit} has old store number {old_store}, "
            "but it is not mapped in OLD_TO_NEW_STORE."
        )
    return new_store


def account_for_description(description: str) -> str:
    account = ACCOUNT_MAP.get(normalize(description))
    if not account:
        raise ValueError(
            f"No QuickBooks account mapping for CustGLDescription '{description}'.\n\n"
            "Add the confirmed QuickBooks account to ACCOUNT_MAP, then convert again."
        )
    return account


def read_xlsx(path: str) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "XLSX support needs openpyxl. Install once with:\n\npy -m pip install openpyxl"
        ) from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(v or "").strip() for v in next(rows)]
    except StopIteration:
        raise ValueError("The Excel file is empty.")
    return [dict(zip(headers, row)) for row in rows if any(v not in (None, "") for v in row)]


def read_csv_input(path: str) -> list[dict]:
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader)


def read_input(path: str) -> list[dict]:
    suffix = Path(path).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx(path)
    if suffix == ".csv":
        return read_csv_input(path)
    raise ValueError("Input must be .xlsx, .xlsm, or .csv")


def find_key(row: dict, requested: str):
    wanted = normalize(requested)
    for key in row.keys():
        if normalize(key) == wanted:
            return key
    return None


def create_rows(input_path: str):
    source = read_input(input_path)
    if not source:
        raise ValueError("The input file contains no invoice rows.")

    required = ["Inv", "Unit", "DueDate", "InvDate", "CustGLDescription", "ExtendedPrice"]
    keymap = {}
    for name in required:
        key = find_key(source[0], name)
        if not key:
            raise ValueError(f"Missing required input column: {name}")
        keymap[name] = key

    total_key = find_key(source[0], "InvTotal")

    # Group by invoice, unit, dates and QuickBooks category account.
    grouped = defaultdict(lambda: Decimal("0.00"))
    expected_totals = {}

    for row in source:
        inv = str(row.get(keymap["Inv"], "")).strip()
        if not inv:
            continue
        # Excel can return whole-number IDs as numbers.
        if inv.endswith(".0"):
            inv = inv[:-2]
        store = store_number(row.get(keymap["Unit"], ""))
        if store not in LOCATION_BY_STORE:
            raise ValueError(f"Store {store} is missing from LOCATION_BY_STORE.")
        if store not in VENDOR_BY_STORE:
            raise ValueError(f"Store {store} is missing from VENDOR_BY_STORE.")

        inv_date = parse_date(row.get(keymap["InvDate"])).strftime("%m/%d/%Y")
        due_date = parse_date(row.get(keymap["DueDate"])).strftime("%m/%d/%Y")
        description = row.get(keymap["CustGLDescription"], "")
        account = account_for_description(description)
        amount = money(row.get(keymap["ExtendedPrice"], 0))

        group_key = (inv, store, inv_date, due_date, account)
        grouped[group_key] += amount

        if total_key:
            expected = money(row.get(total_key, 0))
            expected_totals[(inv, store)] = expected

    if not grouped:
        raise ValueError("No invoice detail rows were found.")

    output = []
    invoice_sums = defaultdict(lambda: Decimal("0.00"))

    for (inv, store, inv_date, due_date, account), amount in sorted(
        grouped.items(), key=lambda item: (item[0][2], item[0][0], normalize(item[0][4]))
    ):
        amount = amount.quantize(CENT, rounding=ROUND_HALF_UP)
        invoice_sums[(inv, store)] += amount
        output.append([
            inv,
            VENDOR_BY_STORE[store],
            "",                         # Mailing Address
            "",                         # Terms
            inv_date,
            due_date,
            LOCATION_BY_STORE[store],
            "",                         # Memo
            "Category Details",
            account,
            "", "", "", "",          # Product/Service, Qty, Rate, Description
            f"{amount:.2f}",
            "", "", "", "",          # Billable, Customer/Project, TaxRate, Class
        ])

    # Validate grouped total against InvTotal when available.
    problems = []
    for key, expected in expected_totals.items():
        actual = invoice_sums.get(key, Decimal("0.00")).quantize(CENT)
        if abs(actual - expected) > CENT:
            problems.append(
                f"Invoice {key[0]} / store {key[1]}: grouped total {actual:.2f} "
                f"does not match InvTotal {expected:.2f}."
            )
    if problems:
        raise ValueError("Conversion stopped because totals do not match:\n\n" + "\n".join(problems))

    return output, invoice_sums


def write_output(rows: list[list], output_path: str):
    with open(output_path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(OUTPUT_HEADERS)
        writer.writerows(rows)


def convert(input_path: str, output_path: str):
    rows, invoice_sums = create_rows(input_path)
    write_output(rows, output_path)
    total = sum(invoice_sums.values(), Decimal("0.00")).quantize(CENT)
    return len(rows), len(invoice_sums), total


class App:
    def __init__(self, root):
        self.root = root
        root.title("PrimeLedgerAI.com - McLane Invoice Converter")
        root.geometry("840x300")
        root.minsize(720, 270)

        frame = ttk.Frame(root, padding=24)
        frame.pack(fill="both", expand=True)

        ttk.Label(
            frame,
            text="McLane Invoice Detail → QuickBooks Online CSV",
            font=("Segoe UI", 18, "bold"),
        ).pack(pady=(0, 18))

        row = ttk.Frame(frame)
        row.pack(fill="x", pady=8)
        ttk.Label(row, text="McLane Input File", width=20).pack(side="left")
        self.input_var = tk.StringVar()
        ttk.Entry(row, textvariable=self.input_var).pack(side="left", fill="x", expand=True)
        ttk.Button(row, text="Browse", command=self.choose_input).pack(side="left", padx=(8, 0))

        ttk.Button(
            frame,
            text="Create QuickBooks CSV",
            command=self.run_conversion,
        ).pack(pady=22, ipadx=22, ipady=7)

        self.status = ttk.Label(
            frame,
            text="Select the McLane XLSX/CSV input file. COA and location mappings are built in.",
        )
        self.status.pack()

    def choose_input(self):
        path = filedialog.askopenfilename(
            title="Select McLane Invoice Detail",
            filetypes=[
                ("Excel files", "*.xlsx *.xlsm"),
                ("CSV files", "*.csv"),
                ("All files", "*.*"),
            ],
        )
        if path:
            self.input_var.set(path)

    def run_conversion(self):
        input_path = self.input_var.get().strip()
        if not input_path:
            messagebox.showwarning("Input required", "Select the McLane invoice detail file.")
            return

        stem = Path(input_path).stem
        output_path = filedialog.asksaveasfilename(
            title="Save QuickBooks CSV",
            initialfile=f"QuickBooks_Ready_{stem}.csv",
            defaultextension=".csv",
            filetypes=[("CSV", "*.csv")],
        )
        if not output_path:
            return

        try:
            row_count, invoice_count, total = convert(input_path, output_path)
            self.status.config(text=f"Created: {Path(output_path).name}")
            messagebox.showinfo(
                "Completed",
                f"QuickBooks CSV created successfully.\n\n"
                f"Invoices: {invoice_count}\n"
                f"Output rows: {row_count}\n"
                f"Total: ${total:,.2f}",
            )
        except Exception as exc:
            messagebox.showerror("Conversion stopped", str(exc))


def main():
    root = tk.Tk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
